# -*- coding: utf-8 -*-
"""W0-2 — 生成/收集 golden 6 类样本到 test/golden/samples/。

设计:
- 真实大样本(c1/c2/火山)**本地产物**,可不上库;缺失时 golden 测试显式 SKIP(见 test/golden-baseline.test.js)。
- 小样本(scan-4p / docx / GBK txt)**可上库**,保证开箱即可跑部分类别。
- 全部可复现:本脚本 + 源路径写在 test/golden/samples/README.md。

类别对照(计划 §W0-2):
  ① 逐字符文字层 ② 正常文字层多图 ③ 纯图扫描件 ④ Office ⑤ 大表格 ⑥ GBK 文本
"""
import shutil
import sys
from pathlib import Path

DEV = Path(r"D:\WorkSpace\Planing-Workdeck\.tmp\dsh-md-convert-dev")
R = Path(r"D:\WorkSpace\Planing-Workdeck\dsh-md-convert-改进研究")
OUT = DEV / "test" / "golden" / "samples"
OUT.mkdir(parents=True, exist_ok=True)
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

# ① 逐字符文字层(Word 导出,逐字符定位) / ⑤ 大表格(13 张真实表)
COPIES = [
    ("char-layer-11p.pdf", R / "corpus" / "c1-brief-11p.pdf", "① 逐字符定位文字层(11p,自研链重建场景)"),
    ("bigtable-34p.pdf", R / "corpus" / "c2-need-34p.pdf", "⑤ 大表格(34p,13 张真实表)"),
    ("textlayer-multi-img.pdf", Path(r"D:\Downloads\火山方舟_Agent 进化_1789355515.pdf"), "② 正常文字层+多图(原 golden 样本)"),
    ("scan-4p.pdf", R / "corpus" / "scan-4p.pdf", "③ 纯图扫描件(4p,路由 OCR)"),
]
for name, src, desc in COPIES:
    dst = OUT / name
    if dst.exists():
        print(f"  已存在 {name}")
        continue
    if not src.exists():
        print(f"  !! 缺源文件,跳过 {name}: {src}")
        continue
    shutil.copy2(src, dst)
    print(f"  + {name} ({dst.stat().st_size / 1024:.0f} KB)  {desc}")

# ④ Office:用 python-docx 生成(含中文标题/列表/表格 → 覆盖 markitdown 原生路径)
docx_path = OUT / "office-cn-table.docx"
if not docx_path.exists():
    try:
        from docx import Document
        from docx.shared import Pt
        d = Document()
        d.add_heading("2026 年泸州老窖窖主节(武汉站)采购需求", level=1)
        d.add_paragraph("本项目共一个标段,公开比选采购一家服务商。服务期自合同签订之日起至 2026-12-31。")
        d.add_heading("一、点位清单", level=2)
        # 注意:此处**不写多余空格**。首版夹具写成 "1 号门广场 玻璃幕墙亮化",
        # 结果基线里 Office 类别 cjk=23.81‰ —— 排查确认那 3 处空格**来自夹具文本自身**,
        # 转换器一处未注入(PDF 类别同指标为 0)。夹具不该带这种噪声,故去掉。
        for it in ["1号门广场:玻璃幕墙亮化", "2-3号门:外围绿化", "地下车库:入口指引"]:
            d.add_paragraph(it, style="List Bullet")
        t = d.add_table(rows=4, cols=3)
        hdr = ["序号", "点位", "内容"]
        for i, h in enumerate(hdr):
            t.cell(0, i).text = h
        for r, row in enumerate([["1", "1号门广场", "玻璃幕墙亮化"], ["2", "2-3号门", "外围绿化"], ["3", "地下车库", "入口指引"]], start=1):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
        d.add_heading("二、最高限价", level=2)
        d.add_paragraph("拦标价 194,000 元(含税)。")  # 中文↔数字之间的空格按规则 R2 保留(不并入归并)
        d.save(str(docx_path))
        print(f"  + {docx_path.name} (生成,含标题/列表/表格)")
    except Exception as e:
        print(f"  !! docx 生成失败(需 python-docx): {e}")

# ⑥ GBK 文本(编码探测路径)
gbk_path = OUT / "gbk-text.txt"
if not gbk_path.exists():
    text = "泸州老窖窖主节(武汉站)采购需求\n一、点位\n1 号门广场:玻璃幕墙亮化\n拦标价 194,000 元(含税)\n"
    gbk_path.write_bytes(text.encode("gbk"))
    print(f"  + {gbk_path.name} (GBK 编码, {gbk_path.stat().st_size} 字节)")

# README:样本来源与恢复方式
readme = OUT / "README.md"
readme.write_text(
    "# golden 样本(W0-2)\n\n"
    "6 类样本,覆盖计划 §W0-2。**大样本是本地资产,不入库**(缺失时 `test:golden` 显式 SKIP 对应类别)。\n\n"
    "| 文件 | 类别 | 来源 | 入库 |\n|---|---|---|---|\n"
    "| `char-layer-11p.pdf` | ① 逐字符文字层 | `改进研究/corpus/c1-brief-11p.pdf` | 否(3.9MB) |\n"
    "| `textlayer-multi-img.pdf` | ② 正常文字层+多图 | `D:\\Downloads\\火山方舟_Agent 进化_1789355515.pdf` | 否(4.6MB) |\n"
    "| `scan-4p.pdf` | ③ 纯图扫描件 | `改进研究/corpus/scan-4p.pdf` | 是 |\n"
    "| `office-cn-table.docx` | ④ Office | `make_golden_samples.py` 生成 | 是 |\n"
    "| `bigtable-34p.pdf` | ⑤ 大表格 | `改进研究/corpus/c2-need-34p.pdf` | 否(27MB) |\n"
    "| `gbk-text.txt` | ⑥ GBK 文本 | `make_golden_samples.py` 生成 | 是 |\n\n"
    "恢复:`python test/make_golden_samples.py`(需源文件在位)。\n"
    "扫描件类别默认不参与 `test:golden`(慢),用 `MDC_GOLDEN_SCAN=1` 打开。\n",
    encoding="utf-8",
)
print(f"  + {readme.name}")
print("\n样本目录内容:")
for p in sorted(OUT.iterdir()):
    if p.is_file():
        print(f"  {p.name:28} {p.stat().st_size / 1024:9.0f} KB")
